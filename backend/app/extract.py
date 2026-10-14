"""Upload processing: read any supported file into pages of paragraphs, suggest its details,
and find the deadlines it sets."""
import csv
import hashlib
import io
import json
import re
from pathlib import Path

import pymupdf
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from . import ocr, store
from .llm import chat
from .rag import block_text

_KINDS = {".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx", ".csv": "xlsx", ".txt": "txt", ".md": "txt",
          ".png": "img", ".jpg": "img", ".jpeg": "img", ".webp": "img"}
_PER_PAGE = 12  # paragraphs per page when the file has no pages of its own
_MIME = {".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
         ".webp": "image/webp", ".txt": "text/plain; charset=utf-8", ".md": "text/plain; charset=utf-8",
         ".csv": "text/csv; charset=utf-8",
         ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
         ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
_MAX_OCR_PAGES = 10
_FRAGMENT = 25  # characters: shorter paragraphs are joined to their neighbour
_CHUNK = 700     # characters: longer paragraphs and tables are split into pieces about this size


def check_supported(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in _KINDS:
        raise ValueError(f"{suffix or 'This'} files aren't supported. Use PDF, Word (.docx), Excel (.xlsx), "
                         "CSV, an image or a text file.")
    return suffix


def extract(filename: str, raw: bytes) -> dict:
    """-> {"kind", "pages": [[block, ...], ...], "ocr": confidence or None, "ocr_engine": name or None}

    Text is read once here; afterwards only the stored paragraphs and embeddings are used.
    """
    suffix = check_supported(filename)
    kind = _KINDS[suffix]
    scans: list[tuple[str, int, str, list[str]]] = []  # OCR results, one per page image
    if kind == "pdf":
        pdf = pymupdf.open(stream=raw, filetype="pdf")
        pages = [_pdf_page(p) for p in pdf]
        if sum(map(len, pages)) == 0:  # no text layer: a scan, so OCR the page images
            scans = [ocr.read(p.get_pixmap(dpi=300).tobytes("png")) for p in list(pdf)[:_MAX_OCR_PAGES]]
            kind, pages = "scan", [_paras(t, keep_lines=True) for t, *_ in scans]
    elif kind == "docx":
        pages = _paginate(_docx(raw))
    elif kind == "xlsx":
        pages = [_rows(raw, suffix)]
    elif kind == "img":
        scans = [ocr.read(raw)]
        pages = _paginate(_paras(scans[0][0], keep_lines=True))
    else:
        pages = _paginate(_paras(raw.decode("utf-8", "ignore")))
    return {
        "kind": kind, "pages": _blocks(pages),
        "ocr": round(sum(s[1] for s in scans) / len(scans)) if scans else None,
        "ocr_engine": "+".join(sorted({s[2] for s in scans})) or None,
        "unsure": sorted({w for s in scans for w in s[3]})[:300],  # for the uploader to check
    }


def _row(cells) -> str:
    """One table row as "cell | cell"; a merged cell (repeated, or empty) appears once."""
    out: list[str] = []
    for c in cells:
        c = " ".join((c or "").split())
        if c and (not out or out[-1] != c):
            out.append(c)
    return " | ".join(out)


def _table(rows) -> str:
    """A table as one paragraph, one line per row, so a date stays next to its subject."""
    return "\n".join(r for r in (_row(row) for row in rows) if r)


def _pdf_page(page) -> list[str]:
    """Paragraphs of a PDF page in reading order. Ruled tables come out row by row; read as plain text,
    many PDFs give a table column by column."""
    tables = page.find_tables().tables
    if not tables:
        return _paras(page.get_text())
    boxes = [pymupdf.Rect(t.bbox) for t in tables]
    parts = [(t.bbox[1], _table(t.extract())) for t in tables]
    parts += [(b[1], " ".join(b[4].split())) for b in page.get_text("blocks")
              if b[6] == 0 and not any(pymupdf.Rect(b[:4]).intersects(box) for box in boxes)]  # text outside tables
    return [text for _, text in sorted(parts, key=lambda p: p[0]) if text.strip()]


def _docx(raw: bytes) -> list[str]:
    """Paragraphs and tables in the order they appear (python-docx's .paragraphs skips tables)."""
    doc = Document(io.BytesIO(raw))
    out = []
    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            out.append(Paragraph(el, doc).text)
        elif tag == "tbl":
            out.append(_table([c.text for c in row.cells] for row in Table(el, doc).rows))
    return out


def keep_original(filename: str, raw: bytes) -> dict:
    """Store the uploaded file (once per content) and return what a document records about it."""
    key = hashlib.sha256(raw).hexdigest()
    mime = _MIME.get(Path(filename).suffix.lower(), "application/octet-stream")
    store.put_file(key, filename, mime, raw)
    info = {"key": key, "name": filename, "mime": mime}
    if mime == "application/pdf":
        info["pages"] = len(pymupdf.open(stream=raw, filetype="pdf"))
    return info


def text_hash(pages: list[list[dict]]) -> str:
    """Fingerprint of a document's text (case and spacing ignored), to spot the same file uploaded twice."""
    text = " ".join(block_text(b) for page in pages for b in page)
    return hashlib.sha256(" ".join(text.lower().split()).encode()).hexdigest() if text.strip() else ""


def _paras(text: str, keep_lines: bool = False) -> list[str]:
    """Paragraphs are separated by blank lines. keep_lines (OCR): a paragraph keeps its line breaks,
    so each row of a scanned table stays one line instead of running into the next."""
    parts = re.split(r"\n\s*\n", text)
    if len(parts) < 3 and not keep_lines:  # no blank lines between paragraphs: fall back to lines
        parts = text.splitlines()
    if keep_lines:
        return ["\n".join(" ".join(ln.split()) for ln in p.splitlines() if ln.strip()) for p in parts if p.strip()]
    return [" ".join(p.split()) for p in parts if p.strip()]


def _paginate(paras: list[str]) -> list[list[str]]:
    paras = [p for p in paras if p.strip()]
    return [paras[i:i + _PER_PAGE] for i in range(0, len(paras), _PER_PAGE)]


def _rows(raw: bytes, suffix: str) -> list[list[str]]:
    if suffix == ".csv":
        rows = csv.reader(io.StringIO(raw.decode("utf-8", "ignore")))
    else:
        rows = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active.iter_rows(values_only=True)
    return [[("" if v is None else str(v)) for v in row] for row in rows if any(v not in (None, "") for v in row)]


def _join_fragments(page: list) -> list:
    """Tiny pieces of text (a stray OCR word, a one-word heading) get no useful search match alone and
    crowd out real paragraphs: attach each to the paragraph that follows it (or, at the end, precedes it)."""
    out, carry = [], []
    for p in page:
        if isinstance(p, str) and len(p.strip()) < _FRAGMENT:
            carry.append(p.strip())
            continue
        if carry and isinstance(p, str):
            p = "\n".join([*carry, p])
            carry = []
        out.append(p)
    if carry:
        if out and isinstance(out[-1], str):
            out[-1] += "\n" + "\n".join(carry)
        else:
            out.append("\n".join(carry))
    return out


def _split_long(text: str) -> list[str]:
    """A long paragraph's meaning gets blurred in search, and one row of a big table gets lost in it.
    Split it into pieces of about _CHUNK characters at line (or sentence) breaks; a table's pieces each
    start with its column headings, so every piece still says what its numbers mean."""
    if len(text) <= _CHUNK * 1.4:
        return [text]
    lines = text.split("\n")
    rows = [ln for ln in lines if " | " in ln]
    header = rows[0] if len(rows) >= 3 else None
    if len(lines) < 3:  # one long line: break at sentences instead
        lines = [s + "." for s in text.split(". ") if s.strip()]
    pieces, cur = [], []
    for ln in lines:
        if cur and len("\n".join(cur)) + len(ln) > _CHUNK:
            pieces.append(cur)
            cur = [header] if header and ln != header else []
        cur.append(ln)
    if cur:
        pieces.append(cur)
    return ["\n".join(p) for p in pieces]


def _blocks(pages: list[list]) -> list[list[dict]]:
    """Give every paragraph (or sheet row) an id that is unique within the document."""
    n, out = 0, []
    split = lambda page: [piece for p in page for piece in (_split_long(p) if isinstance(p, str) else [p])]  # noqa: E731
    for page in (split(_join_fragments(p)) for p in pages):
        blocks = []
        for text in page:
            n += 1
            if isinstance(text, list):
                blocks.append({"id": f"r{n}", "text": text, "kind": "hdr" if n == 1 else "row"})
            else:
                subj = n == 1 and text.lower().startswith(("subject", "sub:", "विषय"))
                blocks.append({"id": f"b{n}", "text": text, "kind": "subj" if subj else "p"})
        if blocks:
            out.append(blocks)
    return out


def suggest(pages: list[list[dict]], filename: str, lists: dict) -> dict:
    """Suggest the document's details, choosing only from the university's own lists."""
    choices = {k: lists.get(k, []) for k in ("TYPES", "AUTHS", "DEPTS", "AUDS", "TAGS", "VIS")}
    text = "\n".join(block_text(b) for page in pages for b in page)[:6000]
    res = chat(
        "You catalogue official documents of Swami Vivekanand Subharti University. Suggest this document's "
        "details. Reply as JSON with keys title, ref (reference or letter number exactly as printed), type, "
        "authority, dept, date, effective, expiry (dates as YYYY-MM-DD), audience (list), tags (list), "
        "visibility, description (one sentence). Pick type from TYPES, authority from AUTHS, dept from DEPTS, "
        "audience from AUDS, tags from TAGS and visibility from VIS. Use \"\" or [] when the text doesn't say.\n"
        + json.dumps(choices, ensure_ascii=False),
        f"File name: {filename}\n\n{text}", as_json=True)
    pick = {"type": "TYPES", "authority": "AUTHS", "dept": "DEPTS", "visibility": "VIS"}
    for k, lst in pick.items():  # drop anything not on the lists
        if res.get(k) not in choices[lst]:
            res[k] = ""
    for k, lst in (("audience", "AUDS"), ("tags", "TAGS")):
        res[k] = [v for v in res.get(k) or [] if v in choices[lst]]
    return {k: v for k, v in res.items() if v not in ("", [], None)}


def deadlines(doc: dict) -> list[dict]:
    """Dated actions the document sets, each pointing at the paragraph it came from."""
    blocks = {(p, b["id"]): b for p, page in enumerate(doc["pages"], 1) for b in page}
    text = "\n".join(f"[page {p}, {bid}] {block_text(b)}" for (p, bid), b in blocks.items())
    res = chat(
        "List every dated action or deadline this document sets. Reply as JSON {\"deadlines\": [{\"title\": "
        "\"short action\", \"date\": \"YYYY-MM-DD\", \"page\": n, \"block\": \"paragraph id\", \"quote\": \"the exact "
        "words from that paragraph that contain the date\", \"owner\": \"who must act\", \"aud\": \"who it "
        "applies to\"}]}. Use only dates written in the text.",
        f"Title: {doc['title']}\n{text}", as_json=True)
    return [d for d in res.get("deadlines") or [] if (d.get("page"), d.get("block")) in blocks and d.get("date")]
