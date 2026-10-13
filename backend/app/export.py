"""Cited answers exported as a Word document."""
import io
import re
from datetime import date

from docx import Document


def _clean(text: str) -> str:
    return re.sub(r"\{(\d+)\}", r"[\1]", text.replace("**", ""))


def answers_docx(title: str, items: list[dict], docs: dict) -> bytes:
    """items: [{q, answer}]. docs: the documents the person may see, by id — citations to anything
    else are shown as restricted."""
    out = Document()
    out.add_heading(title, 1)
    out.add_paragraph(f"SVSU Knowledge Assistant · exported {date.today():%d %b %Y}").italic = True
    for item in items:
        a = item.get("answer") or {}
        out.add_heading(item.get("q", ""), 2)
        if a.get("banner"):
            out.add_paragraph().add_run("Note: " + a["banner"].get("text", "")).italic = True
        for part in a.get("body") or []:
            for key in ("p", "sub"):
                if part.get(key):
                    out.add_paragraph(_clean(part[key]))
            for line in part.get("ul") or []:
                out.add_paragraph(_clean(line), style="List Bullet")
            if t := part.get("table"):
                table = out.add_table(rows=1, cols=len(t["head"]))
                table.style = "Table Grid"
                for cell, text in zip(table.rows[0].cells, t["head"]):
                    cell.text = text
                for row in t["rows"]:
                    for cell, text in zip(table.add_row().cells, row):
                        cell.text = _clean(text)
        if a.get("cites"):
            out.add_paragraph().add_run("Sources").bold = True
            for n, c in enumerate(a["cites"], 1):
                d = docs.get(c.get("doc"))
                src = f"{d['title']} ({d['ref']}), page {c.get('page')}" if d else "A document you can no longer open"
                out.add_paragraph(f"[{n}] {src} — “{c.get('quote', '')}”" if d else f"[{n}] {src}")
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()
